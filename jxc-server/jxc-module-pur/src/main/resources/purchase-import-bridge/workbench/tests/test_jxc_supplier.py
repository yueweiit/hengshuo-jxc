"""非模板识别只能选择原单元格，不能替换业务数值。"""
import hashlib
import pytest
from openpyxl import Workbook
from workbench.jxc_supplier import build_supplier_document, recognition_request
from workbench.jxc_bridge import SourceError


def supplier(tmp_path):
    path = tmp_path / '供应商.xlsx'
    book = Workbook(); sheet = book.active; sheet.title = '订单'
    sheet.append(['PO-1', None, '小规格', '大规格', '报价'])
    sheet.append(['款号', '蓝色', 10, 4, '53.1234567890'])
    sheet['F2'] = '=E2*2'; sheet['G2'] = '=C2*E2'; book.save(path)
    compact = dict(source=dict(sha256=hashlib.sha256(path.read_bytes()).hexdigest()),
                   tables=[dict(source_sheet='订单', headers={'C': '小规格', 'D': '大规格'})],
                   rows=[dict(source_sheet='订单', source_row=2, row_kind='detail', cells={'A': '款号', 'C': '10', 'D': '4'})])
    mapping = dict(format='SupplierMappingV1', fileSha256=compact['source']['sha256'], sheet='订单',
                   needsReview=False, reasons=[], ignoredRows=[], lines=[
                       dict(sourceRow=2, cells=dict(sourceOrderNo='A1', productSpec=['B2','C1'], purchaseQty='C2', purchaseUnitPrice='E2', freight='F2')),
                       dict(sourceRow=2, cells=dict(sourceOrderNo='A1', productSpec=['B2','D1'], purchaseQty='D2', purchaseUnitPrice='E2'))])
    return path, compact, mapping


def test_supplier_columns_split_into_specs_without_changing_source_row_or_prices(tmp_path):
    path, compact, mapping = supplier(tmp_path)
    result = build_supplier_document(path, compact, mapping)
    assert result['source']['parserVersion'] == 'docling-n8n-supplier-v1'
    assert [r['sourceRow'] for r in result['rows']] == [2, 2]
    assert [r['sourceItemKey'] for r in result['rows']] == ['C2', 'D2']
    assert [r['normalizedValues']['purchaseQty'] for r in result['rows']] == ['10.0000', '4.0000']
    assert result['rows'][0]['normalizedValues']['productSpec'] == '蓝色 / 小规格'
    assert result['rows'][0]['normalizedValues']['purchaseUnitPrice'] == '53.1234567890'
    assert result['rows'][0]['normalizedValues']['freight'] is None
    assert result['rows'][0]['formulas']['F'] == '=E2*2'
    assert result['rows'][0]['sourceCells']['purchaseQty'] == 'C2'
    assert result['warnings'][0]['code'] == 'FORMULA_CACHE_MISSING'


@pytest.mark.parametrize('change', ['literal', 'missing_cell', 'duplicate_qty', 'wrong_hash', 'omitted_row', 'needs_review'])
def test_invalid_mapping_preserves_valid_rows_and_reports_each_problem(tmp_path, change):
    path, compact, mapping = supplier(tmp_path)
    if change == 'literal': mapping['lines'][0]['cells']['purchaseQty'] = '999'
    elif change == 'missing_cell': mapping['lines'][0]['cells']['purchaseUnitPrice'] = 'Z99'
    elif change == 'duplicate_qty': mapping['lines'][1]['cells']['purchaseQty'] = 'C2'
    elif change == 'wrong_hash': mapping['fileSha256'] = '0'*64
    elif change == 'omitted_row': compact['rows'].append(dict(source_sheet='订单',source_row=3,row_kind='detail',cells={'A':'遗漏'}))
    else: mapping['needsReview'] = True; mapping['reasons'] = ['单价口径不明确']
    if change == 'wrong_hash':
        with pytest.raises(SourceError): build_supplier_document(path, compact, mapping)
    else:
        result = build_supplier_document(path, compact, mapping)
        assert result['rows']
        assert result['warnings']
        if change in ('literal', 'duplicate_qty'): assert len(result['rows']) == 1
        if change == 'missing_cell': assert result['rows'][0]['normalizedValues']['purchaseUnitPrice'] is None
        if change == 'omitted_row':
            assert any(w['code'] == 'ROW_NOT_STAGED' and w['sourceRow'] == 3 for w in result['warnings'])


def test_model_document_has_no_task_credential_or_arbitrary_path(tmp_path):
    path, compact, _ = supplier(tmp_path)
    compact['credential'] = 'not-model-input'; compact['source']['file'] = 'C:/private/file.xlsx'
    request = recognition_request(compact, None)
    assert request['format'] == 'SupplierDocumentV1'
    assert 'credential' not in str(request)
    assert 'C:/private' not in str(request)


@pytest.mark.parametrize('managed', [False, True])
@pytest.mark.parametrize('needs_review', [False, True])
def test_real_bridge_route_fetches_scoped_source_and_calls_docling_before_n8n(tmp_path, monkeypatch, managed, needs_review):
    from http.server import BaseHTTPRequestHandler, HTTPServer
    from threading import Thread
    from fastapi.testclient import TestClient
    from workbench.jxc_bridge import create_app
    import workbench.documents
    import workbench.jxc_supplier as supplier_module
    import convert
    path, compact, mapping = supplier(tmp_path)
    if needs_review:
        mapping['needsReview'] = True
        mapping['reasons'] = ['规格和采购单号待核对']
        for line in mapping['lines']:
            line['cells'].pop('sourceOrderNo'); line['cells'].pop('productSpec')
    calls=[]
    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *_): pass
        def do_GET(self):
            assert self.headers['X-Purchase-Task']=='101:'+'b'*72
            calls.append('asset');data=path.read_bytes();self.send_response(200);self.end_headers();self.wfile.write(data)
        def do_POST(self):
            assert calls[-1]=='docling'
            if managed:
                assert self.path=='/internal/purchase-import/tasks/11111111-2222-4333-8444-555555555555/fields'
                assert self.headers['X-Purchase-Task']=='101:'+'b'*72
                assert self.headers['Authorization'] is None
            else:
                assert self.path=='/webhook/jxc-purchase-fields-v1'
                assert self.headers['Authorization']=='Bearer fields-key'
            body=__import__('json').loads(self.rfile.read(int(self.headers['Content-Length'])))
            assert 'credential' not in body
            assert any(c['cell']=='A1' and c['value']=='PO-1' for c in body['sourceCells'])
            calls.append('n8n');data=__import__('json').dumps(dict(status='ok',mapping=mapping)).encode()
            self.send_response(200);self.end_headers();self.wfile.write(data)
    server=HTTPServer(('127.0.0.1',0),Handler);thread=Thread(target=server.serve_forever,daemon=True);thread.start()
    root=f'http://127.0.0.1:{server.server_port}'
    monkeypatch.setenv('JXC_SUPPLIER_N8N_URL',root+'/webhook/jxc-purchase-fields-v1');monkeypatch.setenv('JXC_SUPPLIER_N8N_KEY','fields-key')
    monkeypatch.setattr(convert,'convert',lambda *args:calls.append('docling'))
    monkeypatch.setattr(workbench.documents,'prepare_document',lambda *args:compact)
    timeouts=[]
    original_opener=supplier_module.build_opener
    def observed_opener(*handlers):
        opener=original_opener(*handlers)
        class Observed:
            def open(self, *args, **kwargs):
                timeouts.append(kwargs['timeout'])
                return opener.open(*args, **kwargs)
        return Observed()
    monkeypatch.setattr(supplier_module,'build_opener',observed_opener)
    try:
        client=TestClient(create_app(root,'worker-key'))
        body=dict(taskId='11111111-2222-4333-8444-555555555555',sourceId=1,sha256=compact['source']['sha256'],credential='101:'+'b'*72,sheet=None)
        if managed:
            body['managedServices']=True
        response=client.post('/internal/jxc/supplier-parse',headers={'Authorization':'Bearer worker-key'},json=body)
        assert response.status_code==200, response.text
        assert calls==['asset','docling','n8n']
        assert timeouts==[600 if managed else 300]
        document = response.json()['document']
        assert len(document['rows']) == 2
        if needs_review:
            assert all(row['normalizedValues']['sourceOrderNo'] is None and row['normalizedValues']['productSpec'] is None for row in document['rows'])
            assert any(w['code'] == 'SOURCE_REVIEW_REQUIRED' for w in document['warnings'])
    finally:
        server.shutdown();server.server_close();thread.join()


def test_review_flag_and_missing_business_fields_keep_reliable_rows(tmp_path):
    path, compact, mapping = supplier(tmp_path)
    mapping['needsReview'] = True
    mapping['reasons'] = ['单位和规格待核对']
    for line in mapping['lines']:
        line['cells'].pop('sourceOrderNo'); line['cells'].pop('productSpec')
    result = build_supplier_document(path, compact, mapping)
    assert len(result['rows']) == 2
    assert result['rows'][0]['normalizedValues']['sourceOrderNo'] is None
    assert result['rows'][0]['normalizedValues']['productSpec'] is None
    assert result['rows'][0]['normalizedValues']['purchaseQty'] == '10.0000'
    assert result['rows'][0]['sourceCells']['purchaseQty'] == 'C2'
    assert any(w['code'] == 'SOURCE_REVIEW_REQUIRED' for w in result['warnings'])


def test_confirmed_spec_labels_split_original_quantity_cells_without_replacing_prices(tmp_path):
    path, compact, mapping = supplier(tmp_path)
    mapping['needsReview'] = True
    for line in mapping['lines']:
        line['cells'].pop('productSpec')
    overrides = dict(mapping=mapping, specLabels={'C': 'S', 'D': 'M'}, confirm=True)
    result = build_supplier_document(path, compact, mapping, overrides=overrides)
    assert [r['normalizedValues']['productSpec'] for r in result['rows']] == ['S', 'M']
    assert [r['sourceItemKey'] for r in result['rows']] == ['C2', 'D2']
    assert result['rows'][0]['normalizedValues']['purchaseUnitPrice'] == '53.1234567890'
    assert result['rows'][0]['sourceCells']['purchaseQty'] == 'C2'


@pytest.mark.parametrize('changes', [dict(confirm=False), dict(specLabels={'C': ''}), dict(extra='http://localhost'), dict(specLabels={'C2': 'S'})])
def test_page_overrides_require_explicit_confirmation_and_column_labels(tmp_path, changes):
    path, compact, mapping = supplier(tmp_path)
    overrides = dict(mapping=mapping, specLabels={'C': 'S'}, confirm=True)
    overrides.update(changes)
    with pytest.raises(SourceError):
        build_supplier_document(path, compact, mapping, overrides=overrides)


def test_unmapped_original_formulas_are_kept_for_review_not_recalculated(tmp_path):
    path, compact, mapping = supplier(tmp_path)
    result = build_supplier_document(path,compact,mapping)
    row = result['rows'][0]
    assert row['formulas']['G'] == '=C2*E2'
    assert row['cachedValues']['G'] is None
    assert row['rawNumericTexts']['C'] == '10'


def test_original_summary_can_be_explicitly_excluded_without_becoming_purchase_detail(tmp_path):
    from openpyxl import load_workbook
    path, compact, mapping = supplier(tmp_path)
    book = load_workbook(path); book.active.append(['合计', None, 10, 4]); book.save(path)
    compact['source']['sha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
    mapping['fileSha256'] = compact['source']['sha256']
    compact['rows'].append(dict(source_sheet='订单', source_row=3, row_kind='total', cells={'A': '合计'}))
    mapping['ignoredRows'] = [dict(sourceRow=3, reason='合计行，非采购明细')]
    result = build_supplier_document(path, compact, mapping)
    assert [r['sourceRow'] for r in result['rows']] == [2, 2]
    assert any(w['code'] == 'SUPPLIER_ROW_EXCLUDED' and w['sourceRow'] == 3 for w in result['warnings'])
    mapping['ignoredRows'][0]['sourceRow'] = 999
    with pytest.raises(SourceError):
        build_supplier_document(path, compact, mapping)


def test_missing_quantity_stays_unknown_without_losing_other_specs(tmp_path):
    path, compact, mapping = supplier(tmp_path)
    mapping['lines'][0]['cells'].pop('purchaseQty')
    result = build_supplier_document(path, compact, mapping)
    assert len(result['rows']) == 2
    assert result['rows'][0]['normalizedValues']['purchaseQty'] is None
    assert result['rows'][0]['sourceItemKey'] == 'ROW:2'
    assert result['rows'][1]['normalizedValues']['purchaseQty'] == '4.0000'


def test_no_reliable_rows_is_an_explicit_error_with_preserved_diagnostics(tmp_path):
    path, compact, mapping = supplier(tmp_path)
    for line in mapping['lines']: line['cells']['purchaseQty'] = '999'
    with pytest.raises(SourceError) as error:
        build_supplier_document(path, compact, mapping)
    assert error.value.details['warnings']
    assert error.value.details['mapping'] == mapping


def test_source_review_context_keeps_real_title_headers_and_value_examples(tmp_path):
    path, compact, mapping = supplier(tmp_path)
    result = build_supplier_document(path, compact, mapping)
    context = result['sheets'][0]['reviewContext']
    assert context['format'] == 'SourceReviewContextV1'
    assert any(s['kind'] == 'IDENTIFIER' and s['label'] == 'PO-1' for s in context['subjects'])
    assert next(c for c in context['columns'] if c['key'] == 'C')['header'] == '小规格'
    assert next(c for c in context['columns'] if c['key'] == 'C')['examples'] == ['10']
